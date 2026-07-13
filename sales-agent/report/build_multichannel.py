#!/usr/bin/env python3
"""Собирает мультиканальную таблицу аутрича: email + VK + TG с готовыми текстами.
Вход: state/leads.json (+ scratchpad enrichment). Выход: report/outreach-multichannel.xlsx + .csv"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE='sales-agent'
leads=json.load(open(f'{BASE}/state/leads.json'))['leads']

def first_name(contact):
    if not contact: return None
    # «Фамилия Имя Отчество (роль)» или «Имя Отчество, роль»
    core=contact.split('(')[0].split(',')[0].strip()
    parts=core.split()
    if len(parts)>=3: return f'{parts[1]} {parts[2]}'
    if len(parts)==2: return f'{parts[0]} {parts[1]}'
    return core

def city_word(city):
    c=city.split(' (')[0].split('-на-')[0].split(' ')[0]
    return {'Санкт':'Петербург','Ростов':'Ростов','Нижний':'Нижний'}.get(c,c)

def vk_text(l):
    name=first_name(l.get('contact_name'))
    hello=f'{name}, здравствуйте!' if name else 'Здравствуйте!'
    comp=l['company'].split('(')[0].strip()
    if l['niche']=='medicine':
        offer='для одиночных клиник у нас работает канал заявок по фиксированной цене — номер пациента подтверждается по SMS, клиника в Краснодаре так выросла с 252 до 2 270 заявок за год'
    elif l['niche']=='beauty':
        offer='мы ведём лазерную эпиляцию в 19 городах, есть канал с оплатой за заявку — 504 ₽ за подтверждённый контакт'
    else:
        offer='считаем школам бесплатный медиаплан набора к сентябрю — сколько заявок и по какой цене реально получить по вашим районам'
    return (f'{hello} Я Яна, агентство OMG — писала вам на почту {l["email"]} про «{comp}». '
            f'Почта у многих — чёрная дыра, поэтому коротко здесь: {offer}. '
            f'Посчитать бесплатно цифры для вас?')

def tg_text(l):
    name=first_name(l.get('contact_name'))
    hello=f'{name}, здравствуйте!' if name else 'Здравствуйте!'
    comp=l['company'].split('(')[0].strip()
    if l['niche']=='medicine':
        core='заявки на приём по фиксированной цене (SMS-подтверждение каждого номера). Одиночная клиника в Краснодаре выросла в 9 раз по заявкам за год'
    elif l['niche']=='beauty':
        core='заявки на эпиляцию с оплатой за результат — 504 ₽ за подтверждённый контакт (кейс сети в 19 городах)'
    else:
        core='набор к сентябрю: бесплатный медиаплан — сколько заявок реально получить и по какой цене'
    return (f'{hello} Я Яна, агентство OMG — писала вам на почту про «{comp}». '
            f'Суть одной строкой: {core}. Прислать расчёт сюда?')

# отбираем лидов мультиканальной базы: батчи 9-10 + все с phone/vk/tg
rows=[]
for l in leads:
    has_second=l.get('phone') or l.get('lpr_vk') or l.get('lpr_tg') or l.get('vk_group') or l.get('tg_link')
    fresh=l.get('created_at','') >= '2026-07-11'
    if not (fresh or has_second): continue
    d=(l.get('external') or {}).get('draft_id')
    status='отправлено '+l['sent_at'] if l.get('sent_at') else ('черновик в Gmail' if d else 'письмо в outbox')
    rows.append({
        'Компания':l['company'].split('(')[0].strip(),'Город':l['city'].split(' (')[0],'Ниша':{'medicine':'медицина','beauty':'бьюти/лазер','education':'образование'}.get(l['niche'],l['niche']),
        'ЛПР':l.get('contact_name') or '','Email':l['email'],'Тема письма':l.get('subject') or '',
        'Письмо':status,'Дизайн/вариант':l.get('variant') or '','Телефон':l.get('phone') or '',
        'VK-группа':l.get('vk_group') or '','Текст для VK':vk_text(l) if l.get('vk_group') else '',
        'Telegram':l.get('tg_link') or ('поиск по номеру: '+l['phone'] if l.get('phone') else ''),
        'Текст для TG':tg_text(l) if (l.get('tg_link') or l.get('phone')) else '',
        'Правило касаний':'email → D+2 VK/TG → D+3 фоллоу-ап письмом'})

wb=Workbook(); ws=wb.active; ws.title='Мультиканальная база'
headers=list(rows[0].keys()) if rows else []
hf=Font(bold=True,color='FFFFFF'); fill=PatternFill('solid',start_color='C2410C')
for ci,h in enumerate(headers,1):
    c=ws.cell(1,ci,h); c.font=hf; c.fill=fill; c.alignment=Alignment(vertical='center')
for ri,r in enumerate(rows,2):
    for ci,h in enumerate(headers,1):
        c=ws.cell(ri,ci,r[h]); c.alignment=Alignment(vertical='top',wrap_text=(h in('Текст для VK','Текст для TG','ЛПР','Тема письма')))
widths={'Компания':26,'Город':14,'Ниша':12,'ЛПР':30,'Email':28,'Тема письма':34,'Письмо':18,'Дизайн/вариант':14,'Телефон':20,'VK-группа':30,'Текст для VK':60,'Telegram':30,'Текст для TG':60,'Правило касаний':30}
for ci,h in enumerate(headers,1): ws.column_dimensions[get_column_letter(ci)].width=widths.get(h,18)
ws.freeze_panes='A2'
wb.save(f'{BASE}/report/outreach-multichannel.xlsx')

import csv
with open(f'{BASE}/report/outreach-multichannel.csv','w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=headers); w.writeheader(); w.writerows(rows)
print('OK rows:',len(rows))
